"""Excel パーサ（大容量対応・openpyxl read_only ストリーミング）。

売上明細 Excel と取引条件マスタ Excel を読み込み、正規化された辞書に変換する。

大容量ファイル（数百 MB）でも OOM しないよう、pandas ではなく
openpyxl の read_only モードで 1 行ずつストリーミング読み込みする。
また、計算に不要な列は保持せずメモリを最小化する。

両 Excel ともサンプルデータ (2026-05 受領分) のスキーマに合わせている。
"""

from __future__ import annotations

import re
from datetime import date, datetime
from typing import Any, Callable

import openpyxl

# 進捗コールバック: (処理済み行数, 総行数 or None)
ProgressCb = Callable[[int, "int | None"], None]

# 進捗を通知する行間隔
_PROGRESS_EVERY = 2000

_KEY_HEAD_RE = re.compile(r"^\d{10}\d{8}")

_DATE_FORMATS = (
    "%Y/%m/%d",
    "%Y-%m-%d",
    "%Y/%m",
    "%Y-%m",
    "%Y%m%d",
    "%Y年%m月%d日",
    "%Y年%m月",
)


def _is_nan(val: Any) -> bool:
    return isinstance(val, float) and val != val  # NaN 判定


def _to_datetime(val: Any) -> datetime | None:
    """セル値を datetime に変換。read_only+date セルは datetime を返すので大半はそのまま。"""
    if val is None or _is_nan(val):
        return None
    if isinstance(val, datetime):
        return val
    if isinstance(val, date):
        return datetime(val.year, val.month, val.day)
    if isinstance(val, (int, float)):
        # 数値のみでは日付か判別できないため None（read_only は日付セルを datetime で返す）
        return None
    s = str(val).strip()
    if not s:
        return None
    for fmt in _DATE_FORMATS:
        try:
            return datetime.strptime(s, fmt)
        except ValueError:
            continue
    return None


def _to_int(val: Any, default: int = 0) -> int:
    if val is None or _is_nan(val):
        return default
    try:
        return int(val)
    except (ValueError, TypeError):
        try:
            return int(float(val))
        except (ValueError, TypeError):
            return default


def _to_float(val: Any, default: float = 0.0) -> float:
    if val is None or _is_nan(val):
        return default
    try:
        return float(val)
    except (ValueError, TypeError):
        return default


def _to_str(val: Any, default: str | None = None) -> str | None:
    if val is None or _is_nan(val):
        return default
    s = str(val).strip()
    return s if s else default


def make_lookup_key(
    partner_code: int,
    application_month: datetime,
    product: str,
    payment_method: str,
) -> str:
    """マスタ参照用複合キーを生成する。

    形式: {取引先コード10桁}{YYYYMMDD}{商材}{決済方法}
    日付は月初日 (YYYY-MM-01) を YYYYMMDD で連結。
    例: '201208000320260301ずっとPREMIUMプランクレジットカード'
    """
    code_part = str(int(partner_code))
    first_of_month = application_month.replace(day=1)
    date_part = first_of_month.strftime("%Y%m%d")
    product_part = product.strip() if product else ""
    payment_part = payment_method.strip() if payment_method else ""
    return f"{code_part}{date_part}{product_part}{payment_part}"


def _header_index(header_row: tuple) -> dict[str, int]:
    """ヘッダー行から {列名: 最初に出現した列index} のマップを作る。"""
    idx: dict[str, int] = {}
    for i, h in enumerate(header_row):
        if h is None:
            continue
        key = str(h).strip()
        if key and key not in idx:
            idx[key] = i
    return idx


def _cell(row: tuple, idx: dict[str, int], name: str) -> Any:
    i = idx.get(name)
    if i is None or i >= len(row):
        return None
    return row[i]


def parse_sales_excel(
    file_path: str, progress_cb: ProgressCb | None = None
) -> list[dict]:
    """売上明細 Excel をストリーミング読み込みして辞書リストを返す。

    計算に必要な列だけを保持し、複合キーを再生成する（大容量対応）。
    """
    wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
    try:
        ws = wb[wb.sheetnames[0]]
        row_iter = ws.iter_rows(values_only=True)
        try:
            header = next(row_iter)
        except StopIteration:
            return []
        idx = _header_index(header)

        records: list[dict] = []
        count = 0
        for row in row_iter:
            if row is None:
                continue

            record_no = _to_int(_cell(row, idx, "レコードNo（手数料明細用）"))
            partner_code = _to_int(_cell(row, idx, "（Rename）取引先コード"))
            # 完全な空行はスキップ
            if not record_no and not partner_code:
                continue

            application_month = _to_datetime(_cell(row, idx, "申込月"))
            product_name = _to_str(_cell(row, idx, "（Rename）商材"), "") or ""
            payment_method = _to_str(_cell(row, idx, "（Rename）決済方法"), "") or ""

            # 「理由」列に既存の複合キーがあれば優先、なければ再生成
            existing_key = _to_str(_cell(row, idx, "理由"))
            if existing_key and _KEY_HEAD_RE.match(existing_key):
                lookup_key = existing_key
            elif application_month and partner_code:
                lookup_key = make_lookup_key(
                    partner_code, application_month, product_name, payment_method
                )
            else:
                lookup_key = ""

            # 計算に必要な最小限の列のみ保持（メモリ削減）
            records.append(
                {
                    "record_no": record_no,
                    "ファイル区分": _to_str(_cell(row, idx, "ファイル区分")),
                    "partner_code": partner_code,
                    "application_month": application_month,
                    "product_name": product_name,
                    "payment_method": payment_method,
                    "delivery_count": _to_int(_cell(row, idx, "配送個数")),
                    "lookup_key": lookup_key,
                }
            )
            count += 1
            if progress_cb and count % _PROGRESS_EVERY == 0:
                progress_cb(count, None)

        if progress_cb:
            progress_cb(count, count)
        return records
    finally:
        wb.close()


def parse_master_excel(
    file_path: str, progress_cb: ProgressCb | None = None
) -> dict[str, dict]:
    """取引条件マスタ Excel をストリーミング読み込みして複合キー辞書を返す。"""
    wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
    try:
        ws = wb[wb.sheetnames[0]]
        row_iter = ws.iter_rows(values_only=True)
        try:
            header = next(row_iter)
        except StopIteration:
            return {}
        idx = _header_index(header)

        records: dict[str, dict] = {}
        count = 0
        for row in row_iter:
            if row is None:
                continue
            key = _to_str(_cell(row, idx, "キー"))
            if not key:
                continue

            records[key] = {
                "key": key,
                "partner_name": _to_str(_cell(row, idx, "取引先名称"), "") or "",
                "primary_partner_code": _to_int(_cell(row, idx, "一次店コード")),
                "commission_kbn": _to_str(_cell(row, idx, "コミッション区分")),
                "condition_definition": _to_str(_cell(row, idx, "条件適用定義")),
                "payment_definition": _to_str(_cell(row, idx, "支払定義")),
                "product": _to_str(_cell(row, idx, "商材"), "") or "",
                "payment_method": _to_str(_cell(row, idx, "決済方法"), "") or "",
                "basic_commission": _to_float(_cell(row, idx, "基本コミッション")),
                "volume_incentive": _to_float(_cell(row, idx, "ボリュームインセンティブ")),
                "special_commission_1": _to_float(_cell(row, idx, "特別コミッション")),
                "special_commission_2": _to_float(_cell(row, idx, "特別コミッション②")),
                "qi_scope": _to_float(_cell(row, idx, "QI適用範囲")),
                "qi_split_period": _to_float(_cell(row, idx, "QI分割計上期間")),
                "debit_initial_fee": _to_float(_cell(row, idx, "口振分割時初回手数料")),
                "referral_kbn": _to_float(_cell(row, idx, "紹介制度区分")),
                "referral_commission": _to_float(_cell(row, idx, "紹介制度コミッション")),
                "continuous_flag_25_37": _to_float(
                    _cell(row, idx, "25・37ヶ月目以降継続コミッションフラグ")
                ),
                "continuous_commission": _to_float(_cell(row, idx, "継続コミッション")),
                "pap_kbn": _to_str(_cell(row, idx, "PAP区分")),
                "pap_commission": _to_float(_cell(row, idx, "PAPコミッション")),
                "pas_kbn": _to_str(_cell(row, idx, "PAS区分")),
                "pas_commission": _to_float(_cell(row, idx, "PASコミッション")),
                "ph_kbn": _to_str(_cell(row, idx, "デリキチ区分")),
                "ph_commission": _to_float(_cell(row, idx, "デリキチコミッション"))
                + _to_float(_cell(row, idx, "6Lコミッション")),
                "return_condition": _to_str(_cell(row, idx, "戻入条件")),
                "return_full_condition": _to_str(_cell(row, idx, "戻入全額条件")),
                "return_half_condition": _to_str(_cell(row, idx, "戻入半額条件")),
                "penalty": _to_float(_cell(row, idx, "違約金")),
            }
            count += 1
            if progress_cb and count % _PROGRESS_EVERY == 0:
                progress_cb(count, None)

        if progress_cb:
            progress_cb(count, count)
        return records
    finally:
        wb.close()
